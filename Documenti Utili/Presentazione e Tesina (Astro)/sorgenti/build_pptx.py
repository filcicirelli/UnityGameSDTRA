# -*- coding: utf-8 -*-
"""Genera 'Presentazione_Astro.pptx' — Politecnico di Bari. Diagrammi a blocchi nativi."""
import os, json
HERE = os.path.dirname(os.path.abspath(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

OUT = os.environ.get("OUT_PPTX", os.path.join(HERE, "Presentazione_Astro.pptx"))
LOGO = os.path.join(HERE, "logo_polibari.png")
LIT_PATH = os.path.join(HERE, "lit_ppt.json")

if os.path.exists(LIT_PATH):
    with open(LIT_PATH, encoding="utf-8") as f:
        LIT = json.load(f)
else:
    LIT = {
        "tipoRiabilitazione": "Riabilitazione motoria dell'arto superiore e della coordinazione occhio-mano (placeholder).",
        "vantaggi": ["Motivazione e aderenza", "Ripetizione ad alta intensità", "Feedback immediato", "Adattabilità"],
        "evidenze": [{"punto": "I serious games aumentano la motivazione (placeholder).", "riferimento": "(Autore, anno)"}],
        "risultatiAttesi": ["Maggiore tempo di esercizio", "Migliore coordinazione"],
        "brevePerSlide": ["Più motivazione → più ripetizioni", "Feedback immediato → apprendimento motorio",
                          "Adattabile e a basso costo", "Utilizzabile anche a casa"],
        "bibliografia": [{"citazione": "(placeholder citazione)", "nota": ""}],
    }

# ---------- Palette (brand Politecnico di Bari) ----------
TEAL      = RGBColor(0x00, 0x9A, 0x9A)
TEAL_DARK = RGBColor(0x00, 0x6E, 0x6C)
TEAL_DEEP = RGBColor(0x00, 0x4C, 0x4B)
TEAL_MID  = RGBColor(0x3E, 0xAE, 0xAC)
TEAL_LT   = RGBColor(0xE3, 0xF3, 0xF2)
TEAL_LT2  = RGBColor(0xCF, 0xEA, 0xE8)
AMBER     = RGBColor(0xEF, 0x9E, 0x16)
AMBER_INK = RGBColor(0x4A, 0x33, 0x00)
ORANGE    = RGBColor(0xD9, 0x66, 0x27)
REDT      = RGBColor(0xCB, 0x49, 0x3E)
RED_LT    = RGBColor(0xF7, 0xE3, 0xE0)
GREEN     = RGBColor(0x3C, 0x92, 0x4B)
GREEN_LT  = RGBColor(0xE3, 0xF1, 0xE4)
INK       = RGBColor(0x21, 0x2B, 0x2B)
GREY      = RGBColor(0x5E, 0x6E, 0x6E)
GREY_LT   = RGBColor(0xCB, 0xD5, 0xD5)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
PANEL     = RGBColor(0xF3, 0xF8, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
_page = {"n": 0}

# ---------- Helpers ----------
def blank():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    return s

def rect(s, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, line_w=Pt(1.0)):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = line_w
    return sh

def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    b = s.shapes.add_textbox(x, y, w, h)
    tf = b.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    return b, tf

def _apply(run, o):
    f = run.font
    f.size = Pt(o.get("size", 16)); f.bold = o.get("bold", False)
    f.italic = o.get("italic", False); f.name = o.get("font", "Calibri")
    f.color.rgb = o.get("color", INK)

def para(tf, runs, align=PP_ALIGN.LEFT, sa=4, sb=0, ls=None, first=False, bullet=False, bcolor=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(sa); p.space_before = Pt(sb)
    if ls: p.line_spacing = ls
    if bullet:
        rb = p.add_run(); rb.text = "▪  "
        _apply(rb, dict(size=runs[0][1].get("size", 16), bold=True, color=bcolor or TEAL))
    for t, o in runs:
        segs = t.split("\n")
        for i, seg in enumerate(segs):
            if i > 0:
                br = p._p.makeelement(qn("a:br"), {}); p._p.append(br)
            r = p.add_run(); r.text = seg; _apply(r, o)
    return p

def node(s, x, y, w, h, title, sub=None, fill=TEAL, txt=WHITE, ts=14, ss=10.5,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, bold=True, align=PP_ALIGN.CENTER, sub_txt=None):
    sh = rect(s, x, y, w, h, fill, shape=shape, line=line)
    try:
        sh.adjustments[0] = 0.09
    except Exception:
        pass
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.05); tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.01); tf.margin_bottom = Inches(0.01)
    para(tf, [(title, dict(size=ts, bold=bold, color=txt))], align=align, sa=0, first=True, ls=1.0)
    if sub:
        para(tf, [(sub, dict(size=ss, color=sub_txt or txt))], align=align, sa=0, sb=1, ls=1.0)
    return sh

def arrow(s, x1, y1, x2, y2, color=TEAL_DARK, w=Pt(2.0), dash=None, end=True, begin=False):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cn.shadow.inherit = False
    cn.line.color.rgb = color; cn.line.width = w
    ln = cn.line._get_or_add_ln()
    if dash:
        ln.append(ln.makeelement(qn("a:prstDash"), {"val": dash}))
    if begin:
        ln.append(ln.makeelement(qn("a:headEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    if end:
        ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return cn

def label(s, x, y, w, text, size=10, color=TEAL_DARK, bold=True, align=PP_ALIGN.CENTER):
    _, tf = tb(s, x, y, w, Inches(0.26), anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(text, dict(size=size, bold=bold, color=color))], align=align, sa=0, first=True)

def add_logo(s):
    s.shapes.add_picture(LOGO, Inches(11.30), Inches(0.30), height=Inches(0.40))

def footer(s, dark=False):
    _page["n"] += 1
    rect(s, Inches(0.55), Inches(7.04), Inches(12.23), Pt(1.2), GREY_LT)
    _, tf = tb(s, Inches(0.55), Inches(7.10), Inches(10), Inches(0.3))
    para(tf, [("Astro · Videogioco per la riabilitazione motoria", dict(size=9, color=GREY))], first=True, sa=0)
    _, tf2 = tb(s, Inches(11.8), Inches(7.10), Inches(0.98), Inches(0.3))
    para(tf2, [(str(_page["n"]), dict(size=9, color=GREY, bold=True))], align=PP_ALIGN.RIGHT, first=True, sa=0)

def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(0.16), TEAL)
    add_logo(s)
    _, tf = tb(s, Inches(0.55), Inches(0.40), Inches(9.5), Inches(0.3))
    para(tf, [(kicker.upper(), dict(size=12, bold=True, color=TEAL))], first=True, sa=0)
    _, tf2 = tb(s, Inches(0.55), Inches(0.70), Inches(10.6), Inches(0.95))
    para(tf2, [(title, dict(size=27, bold=True, color=TEAL_DEEP))], first=True, sa=0, ls=1.0)
    footer(s)
    return s

def notes(s, text):
    s.notes_slide.notes_text_frame.text = text

def chip(s, x, y, w, h, key, val, kcolor=TEAL_DARK):
    rect(s, x, y, w, h, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
    _, tf = tb(s, x + Inches(0.12), y, w - Inches(0.2), h, anchor=MSO_ANCHOR.MIDDLE)
    para(tf, [(key + ":  ", dict(size=11, bold=True, color=kcolor)), (val, dict(size=11, color=INK))],
         first=True, sa=0, ls=1.0)

# =========================================================================
# 1 — COPERTINA
# =========================================================================
def slide_cover():
    s = blank()
    rect(s, 0, 0, Inches(4.6), SH, TEAL_DEEP)
    rect(s, Inches(4.6), 0, Inches(0.10), SH, AMBER)
    rect(s, 0, 0, Inches(4.6), Inches(0.16), AMBER)
    for cx, cy, d, col in [(1.0, 5.7, 0.5, TEAL_MID), (3.55, 1.15, 0.7, TEAL),
                           (2.0, 6.5, 0.30, AMBER), (3.9, 5.1, 0.22, WHITE), (0.7, 4.9, 0.16, TEAL_MID)]:
        rect(s, Inches(cx), Inches(cy), Inches(d), Inches(d), col, shape=MSO_SHAPE.OVAL)
    node(s, Inches(0.75), Inches(2.45), Inches(2.0), Inches(2.0), "★", fill=AMBER, ts=70, shape=MSO_SHAPE.OVAL)
    _, tf = tb(s, Inches(0.55), Inches(4.70), Inches(3.7), Inches(1.2))
    para(tf, [("ASTRO", dict(size=42, bold=True, color=WHITE))], first=True, sa=0)
    para(tf, [("l'alieno esploratore", dict(size=16, italic=True, color=TEAL_LT))], sa=0)
    _, tfp = tb(s, Inches(0.55), Inches(6.45), Inches(3.7), Inches(0.7))
    para(tfp, [("POLITECNICO DI BARI", dict(size=13, bold=True, color=WHITE))], first=True, sa=1)
    para(tfp, [("Corso di Programmazione / Informatica", dict(size=10, color=TEAL_LT))], sa=0)

    add_logo(s)
    _, tt = tb(s, Inches(5.15), Inches(1.55), Inches(7.6), Inches(2.7))
    para(tt, [("Un videogioco per la", dict(size=24, color=GREY))], first=True, sa=2)
    para(tt, [("riabilitazione motoria", dict(size=40, bold=True, color=TEAL_DEEP))], sa=4, ls=1.0)
    para(tt, [("Progettazione di un'interfaccia software che supporta il paziente durante l'esercizio riabilitativo, "
               "trasformandolo in gioco.", dict(size=15.5, color=GREY))], sb=6, ls=1.12)
    meta = [("Studente", "________________________"), ("Relatore", "________________________"),
            ("Corso", "________________________"), ("Anno Accademico", "20__ / 20__")]
    for i, (k, v) in enumerate(meta):
        x = 5.15 + (i % 2) * 3.95; y = 4.55 + (i // 2) * 0.98
        rect(s, Inches(x), Inches(y), Inches(3.7), Inches(0.80), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
        _, tfm = tb(s, Inches(x + 0.18), Inches(y + 0.11), Inches(3.4), Inches(0.6))
        para(tfm, [(k.upper(), dict(size=9.5, bold=True, color=TEAL))], first=True, sa=1)
        para(tfm, [(v, dict(size=13, color=INK))], sa=0)
    notes(s, "Apertura. Presentati e di' in una frase: «Ho sviluppato da zero un videogioco che aiuta un paziente a fare "
              "esercizi di riabilitazione motoria, controllandolo anche con i movimenti del braccio davanti alla webcam.»")

# =========================================================================
# 2 — INDICE
# =========================================================================
def slide_indice():
    s = blank()
    rect(s, 0, 0, SW, Inches(0.16), TEAL); add_logo(s)
    _, tf = tb(s, Inches(0.55), Inches(0.52), Inches(9), Inches(0.9))
    para(tf, [("Di cosa parliamo", dict(size=30, bold=True, color=TEAL_DEEP))], first=True, sa=0)
    items = [("01", "Il progetto e l'obiettivo", "Cosa ho costruito e perché"),
             ("02", "Come funziona il gioco", "Concept e flusso di gioco (schema a blocchi)"),
             ("03", "L'architettura del software", "I moduli e come dialogano (schema a blocchi)"),
             ("04", "Le tre modalità di comando", "Mouse, dito (webcam), joystick"),
             ("05", "Il feedback al paziente", "Risposta sonora e visiva immediata"),
             ("06", "Livelli e parametri regolabili", "Difficoltà e accessibilità su misura"),
             ("07", "La riabilitazione e le evidenze", "Tipo di terapia, vantaggi, letteratura"),
             ("08", "Conclusioni e sviluppi futuri", "Cosa funziona e dove si può arrivare")]
    for i, (n, t, d) in enumerate(items):
        col = i // 4; row = i % 4
        x = 0.7 + col * 6.3; y = 1.62 + row * 1.28
        node(s, Inches(x), Inches(y), Inches(0.95), Inches(0.95), n, fill=TEAL if col == 0 else TEAL_DARK, ts=22)
        _, tfi = tb(s, Inches(x + 1.15), Inches(y + 0.05), Inches(4.85), Inches(1.05), anchor=MSO_ANCHOR.MIDDLE)
        para(tfi, [(t, dict(size=16, bold=True, color=INK))], first=True, sa=1)
        para(tfi, [(d, dict(size=11.5, color=GREY))], sa=0)
    footer(s)
    notes(s, "Filo: prima 'come l'ho fatto' (gioco + software), poi 'a cosa serve' (riabilitazione ed evidenze).")

# =========================================================================
# 3 — OBIETTIVO
# =========================================================================
def slide_obiettivo():
    s = header(blank(), "Il progetto", "Trasformare l'esercizio in un gioco")
    _, tf = tb(s, Inches(0.55), Inches(1.80), Inches(6.7), Inches(5.0))
    blocks = [("Il problema. ", "la riabilitazione motoria richiede di ripetere molte volte gli stessi movimenti: "
               "è efficace ma noioso, e la motivazione cala — soprattutto nei bambini."),
              ("L'idea. ", "trasformare l'esercizio in un gioco spaziale. Il paziente muove «Astro» per raccogliere "
               "stelline ed evitare ostacoli: continua a muoversi divertendosi."),
              ("Il cuore del progetto. ", "il movimento si può fare con il braccio nello spazio, inquadrato dalla "
               "webcam: così si allenano gesti ampi e controllati."),
              ("Obiettivo del lavoro. ", "dimostrare competenze di programmazione costruendo da zero un'interfaccia "
               "software realmente utile in ambito riabilitativo.")]
    first = True
    for head, rest in blocks:
        para(tf, [(head, dict(size=14.5, bold=True, color=TEAL_DARK)), (rest, dict(size=14.5, color=INK))],
             first=first, sa=11, bullet=True, ls=1.1); first = False
    rect(s, Inches(7.55), Inches(1.80), Inches(5.25), Inches(4.95), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
    _, tfh = tb(s, Inches(7.85), Inches(2.05), Inches(4.7), Inches(0.5))
    para(tfh, [("In sintesi", dict(size=15, bold=True, color=TEAL_DARK))], first=True, sa=0)
    feats = [("Sviluppato in Unity", "linguaggio C#"),
             ("Tutto generato da codice", "nessun file grafico o audio esterno"),
             ("3 modalità di comando", "mouse · dito (webcam) · joystick"),
             ("Feedback multisensoriale", "suono + reazione visiva di Astro"),
             ("Parametri regolabili", "difficoltà e accessibilità su misura"),
             ("Multipiattaforma", "funziona su Windows e Mac")]
    yy = 2.60
    for h, d in feats:
        node(s, Inches(7.90), Inches(yy + 0.03), Inches(0.28), Inches(0.28), "★", fill=AMBER, ts=11, shape=MSO_SHAPE.OVAL)
        _, tff = tb(s, Inches(8.32), Inches(yy - 0.05), Inches(4.25), Inches(0.72), anchor=MSO_ANCHOR.MIDDLE)
        para(tff, [(h, dict(size=13, bold=True, color=INK)), (" — " + d, dict(size=12.5, color=GREY))], first=True, sa=0, ls=1.04)
        yy += 0.69
    notes(s, "Messaggio: non è 'solo un gioco', è un'interfaccia per la terapia. Tutto scritto da codice (competenza di "
              "programmazione); la webcam allena il movimento del braccio.")

# =========================================================================
# 4 — CONCEPT
# =========================================================================
def slide_concept():
    s = header(blank(), "Come funziona il gioco", "Il concept in pochi secondi")
    cards = [("MUOVI ASTRO", "Astro segue il puntatore: mouse, dito davanti alla webcam o joystick. Dove punti, lui va.", TEAL),
             ("RACCOGLI LE STELLINE", "Passa vicino alle stelline per raccoglierle tutte: è la Fase 1 di ogni livello.", AMBER),
             ("EVITA GLI OSTACOLI", "Asteroidi e bombe fanno perdere una vita: hai 5 vite e 60 secondi per livello.", REDT),
             ("RAGGIUNGI LA PORTA", "Raccolte tutte le stelline appare la porta: raggiungila e passi al livello dopo.", GREEN)]
    x0, w, gap = 0.55, 2.96, 0.135
    for i, (t, d, col) in enumerate(cards):
        x = x0 + i * (w + gap)
        rect(s, Inches(x), Inches(1.95), Inches(w), Inches(3.15), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=GREY_LT)
        rect(s, Inches(x), Inches(1.95), Inches(w), Inches(0.62), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        _, tf = tb(s, Inches(x + 0.16), Inches(2.04), Inches(w - 0.32), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(t, dict(size=13, bold=True, color=WHITE))], align=PP_ALIGN.CENTER, first=True, sa=0)
        node(s, Inches(x + w / 2 - 0.42), Inches(2.82), Inches(0.84), Inches(0.84), str(i + 1), fill=col, ts=30, shape=MSO_SHAPE.OVAL)
        _, tf2 = tb(s, Inches(x + 0.22), Inches(3.85), Inches(w - 0.44), Inches(1.15))
        para(tf2, [(d, dict(size=12, color=INK))], align=PP_ALIGN.CENTER, first=True, sa=0, ls=1.06)
    rect(s, Inches(0.55), Inches(5.40), Inches(12.23), Inches(1.22), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tfb = tb(s, Inches(0.95), Inches(5.56), Inches(11.45), Inches(0.95), anchor=MSO_ANCHOR.MIDDLE)
    para(tfb, [("Tre livelli a difficoltà crescente.  ", dict(size=14, bold=True, color=TEAL_DARK)),
               ("Ogni azione — giusta o sbagliata — riceve subito una risposta sonora e visiva: il paziente capisce al "
                "volo se il movimento è corretto.", dict(size=14, color=INK))], first=True, sa=0, ls=1.06)
    notes(s, "Il bambino non 'fa fisioterapia': gioca. La terapia è nel movimento che fa per giocare. "
              "Se puoi, mostra una schermata reale del gioco.")

# =========================================================================
# 5 — FLUSSO DI GIOCO (schema a blocchi 1)
# =========================================================================
def slide_flusso():
    s = header(blank(), "Schema a blocchi — 1", "Come si svolge una partita")
    y, h, w = 1.95, 1.00, 2.16
    xs = [0.55, 3.05, 5.55, 8.05, 10.55]
    spine = [("1 · SCHERMATA\nINIZIALE", "Scegli il comando:\nmouse / dito / joystick", TEAL_DARK),
             ("2 · PRONTI…", "1,5 s di grazia,\npoi parte il tempo", TEAL),
             ("3 · FASE 1", "Raccogli tutte\nle stelline", TEAL),
             ("4 · FASE 2", "Raggiungi la\nporta che appare", TEAL),
             ("5 · LIVELLO\nCOMPLETATO", "Si passa al\nlivello successivo", GREEN)]
    cy = y + h / 2
    for i, (t, sub, col) in enumerate(spine):
        node(s, Inches(xs[i]), Inches(y), Inches(w), Inches(h), t, sub=sub, fill=col, ts=12.5, ss=10)
        if i > 0:
            arrow(s, Inches(xs[i] - 0.34), Inches(cy), Inches(xs[i]), Inches(cy))
    label(s, Inches(2.50), Inches(1.56), Inches(1.1), "GIOCA", color=ORANGE)
    label(s, Inches(7.45), Inches(1.56), Inches(1.25), "tutte raccolte", color=TEAL_DARK, size=9)
    # vittoria finale
    vy = 3.55
    node(s, Inches(10.55), Inches(vy), Inches(2.16), Inches(0.82), "VITTORIA FINALE", sub="dopo il livello 3",
         fill=AMBER, txt=AMBER_INK, ts=13, ss=9.5)
    arrow(s, Inches(10.55 + w / 2), Inches(y + h), Inches(10.55 + w / 2), Inches(vy), color=AMBER)
    # loop prossimo livello
    ry = 3.32
    arrow(s, Inches(xs[2] + w / 2), Inches(y + h), Inches(xs[2] + w / 2), Inches(ry), color=TEAL_MID, end=False)
    arrow(s, Inches(xs[1] + w / 2), Inches(ry), Inches(xs[2] + w / 2), Inches(ry), color=TEAL_MID, dash="dash", begin=True, end=False)
    arrow(s, Inches(xs[1] + w / 2), Inches(ry), Inches(xs[1] + w / 2), Inches(y + h), color=TEAL_MID)
    label(s, Inches(3.55), Inches(3.10), Inches(3.4), "PROSSIMO LIVELLO  (1 → 2 → 3)", color=TEAL_MID, size=9.5)
    # corsia errori
    ey, eh, ew = 4.92, 0.95, 2.78
    exs = [0.55, 3.75, 6.95, 10.15]
    rect(s, Inches(0.34), Inches(ey - 0.34), Inches(12.66), Inches(1.66), RED_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    label(s, Inches(0.55), Inches(ey - 0.32), Inches(4.5), "SE SBAGLI (in qualunque momento)", color=REDT, align=PP_ALIGN.LEFT, size=10)
    err = [("ERRORE", "asteroide, bomba\noppure tempo scaduto", REDT),
           ("REAZIONE", "Astro si schiaccia (rosso)\n+ suono basso e gentile", RGBColor(0xC2, 0x55, 0x4B)),
           ("− 1 VITA  ♥", "punteggio può calare;\nbreve invulnerabilità", RGBColor(0xC2, 0x55, 0x4B)),
           ("VITE = 0 ?", "Sì → GAME OVER (Riprova)\nNo → continui a giocare", RGBColor(0xAE, 0x42, 0x39))]
    ecy = ey + eh / 2
    for i, (t, sub, col) in enumerate(err):
        node(s, Inches(exs[i]), Inches(ey), Inches(ew), Inches(eh), t, sub=sub, fill=col, ts=12.5, ss=9.5)
        if i > 0:
            arrow(s, Inches(exs[i] - 0.42), Inches(ecy), Inches(exs[i]), Inches(ecy), color=REDT)
    arrow(s, Inches(6.6), Inches(y + h), Inches(6.6), Inches(ey - 0.34), color=REDT, dash="dash")
    notes(s, "Primo schema a blocchi: il ciclo di gioco. Due fasi per livello (raccogli stelline → raggiungi la porta), "
              "tre livelli in sequenza, e una 'corsia' rossa per gli errori (vite, game over). Dopo un errore c'è una breve "
              "invulnerabilità per non perdere tutte le vite di colpo.")

# =========================================================================
# 6 — ARCHITETTURA (schema a blocchi 2)
# =========================================================================
def slide_architettura():
    s = header(blank(), "Schema a blocchi — 2", "Com'è fatto il software (i moduli)")
    node(s, Inches(0.55), Inches(1.66), Inches(12.23), Inches(0.50),
         "AVVIO  —  all'accensione prepara la scena (telecamera) e crea i moduli principali",
         fill=TEAL_DEEP, ts=13)
    gx, gy, gw, gh = 4.78, 2.92, 3.78, 1.42
    node(s, Inches(gx), Inches(gy), Inches(gw), Inches(gh), "GESTORE GIOCO",
         sub="il «cervello»: stato, punteggio,\nvite, tempo, livello, eventi", fill=TEAL_DARK, ts=16, ss=11)
    arrow(s, Inches(gx + gw / 2), Inches(2.16), Inches(gx + gw / 2), Inches(gy), color=TEAL_DARK)
    node(s, Inches(0.55), Inches(2.55), Inches(3.5), Inches(0.95), "COMANDI",
         sub="Mouse · Dito (webcam) · Joystick\n→ «dove punta il paziente?»", fill=TEAL, ts=14, ss=10)
    node(s, Inches(0.55), Inches(3.78), Inches(3.5), Inches(0.85), "ASTRO",
         sub="segue il puntatore; tocca\nstelline / ostacoli / porta", fill=TEAL_MID, ts=14, ss=10)
    arrow(s, Inches(2.30), Inches(3.50), Inches(2.30), Inches(3.78))
    arrow(s, Inches(4.05), Inches(4.18), Inches(gx), Inches(gy + gh * 0.62), color=TEAL_DARK)
    label(s, Inches(3.55), Inches(4.28), Inches(1.4), "eventi", color=TEAL_DARK, size=9)
    node(s, Inches(gx), Inches(4.66), Inches(gw), Inches(0.80), "INTERFACCIA GIOCO",
         sub="HUD a schermo, cursore, pulsanti\naccessibili 'a permanenza' (dwell)", fill=TEAL, ts=13.5, ss=9.5)
    arrow(s, Inches(gx + gw / 2), Inches(gy + gh), Inches(gx + gw / 2), Inches(4.66), color=TEAL_DARK)
    rx, rw, rh = 9.32, 3.46, 0.78
    right = [("CARICATORE LIVELLI", "+ Definizione livelli (i 3 livelli)"),
             ("OGGETTI DEL GIOCO", "stelline, asteroidi, bombe, porta"),
             ("FEEDBACK PAZIENTE", "+ Fabbrica suoni (audio da codice)"),
             ("FABBRICA IMMAGINI", "tutte le grafiche disegnate da codice")]
    for (t, sub), yy in zip(right, [2.55, 3.50, 4.45, 5.40]):
        node(s, Inches(rx), Inches(yy), Inches(rw), Inches(rh), t, sub=sub, fill=TEAL, ts=12.5, ss=9.5)
        arrow(s, Inches(gx + gw), Inches(gy + gh / 2), Inches(rx), Inches(yy + rh / 2), color=TEAL_DARK, w=Pt(1.4))
    node(s, Inches(0.55), Inches(6.32), Inches(12.23), Inches(0.55),
         "PAGINE DEI PARAMETRI (solo valori, niente logica):  Impostazioni · ParametriComandi · ParametriWebcam · ParametriFeedback",
         fill=AMBER, txt=AMBER_INK, ts=12.5)
    notes(s, "Secondo schema a blocchi: l'architettura. Concetto chiave: ogni modulo ha un compito solo (separazione delle "
              "responsabilità). 'Comandi' è un punto unico: Astro chiede solo 'dove punto?', non sa se rispondi col mouse, la "
              "webcam o il joystick — per questo aggiungere una modalità non tocca il resto. Banda gialla: tutti i numeri "
              "regolabili sono raccolti in 'pagine dei parametri', separate dalla logica.")

# =========================================================================
# 7 — MODALITA' DI COMANDO
# =========================================================================
def slide_comandi():
    s = header(blank(), "Le tre modalità di comando", "Lo stesso gioco, adattato al paziente")
    modes = [("MOUSE", "Controllo fine del polso", "Per chi ha già buona precisione. Sempre disponibile come riserva.", TEAL_DARK),
             ("DITO (WEBCAM)", "Movimento ampio del braccio", "Si muove un evidenziatore fluo davanti alla webcam: allena gesti ampi e controllati nello spazio.", TEAL),
             ("JOYSTICK", "Leva, frecce o tasti WASD", "Per chi usa già ausili con leva o preferisce un comando fisico.", AMBER)]
    x0, w, gap = 0.55, 3.71, 0.55
    for i, (t, st, d, col) in enumerate(modes):
        x = x0 + i * (w + gap)
        rect(s, Inches(x), Inches(1.82), Inches(w), Inches(2.62), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=GREY_LT)
        rect(s, Inches(x), Inches(1.82), Inches(w), Inches(0.74), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        _, tf = tb(s, Inches(x + 0.2), Inches(1.91), Inches(w - 0.4), Inches(0.55), anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(t, dict(size=17, bold=True, color=WHITE))], align=PP_ALIGN.CENTER, first=True, sa=0)
        _, tf2 = tb(s, Inches(x + 0.25), Inches(2.78), Inches(w - 0.5), Inches(1.55))
        para(tf2, [(st, dict(size=13.5, bold=True, color=(ORANGE if col == AMBER else col)))], first=True, align=PP_ALIGN.CENTER, sa=6)
        para(tf2, [(d, dict(size=12.5, color=INK))], align=PP_ALIGN.CENTER, sa=0, ls=1.06)
    rect(s, Inches(0.55), Inches(4.92), Inches(12.23), Inches(1.02), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    arrow(s, Inches(2.40), Inches(4.44), Inches(5.95), Inches(4.92), color=TEAL_DARK)
    arrow(s, Inches(6.65), Inches(4.44), Inches(6.65), Inches(4.92), color=TEAL_DARK)
    arrow(s, Inches(10.90), Inches(4.44), Inches(7.35), Inches(4.92), color=TEAL_DARK)
    _, tfh = tb(s, Inches(0.9), Inches(5.06), Inches(11.5), Inches(0.78), anchor=MSO_ANCHOR.MIDDLE)
    para(tfh, [("Un punto unico decide «dove punta il paziente» → Astro segue.  ", dict(size=14, bold=True, color=TEAL_DARK)),
               ("Cambiare comando non cambia il resto del gioco; se la webcam manca, si torna automaticamente al mouse.",
                dict(size=14, color=INK))], first=True, sa=0, ls=1.06)
    _, tfr = tb(s, Inches(0.55), Inches(6.14), Inches(12.2), Inches(0.7))
    para(tfr, [("Perché conta per la terapia:  ", dict(size=12.5, bold=True, color=ORANGE)),
               ("si sceglie il comando in base alle capacità del paziente e all'obiettivo motorio — controllo fine col "
                "mouse, movimenti ampi del braccio con la webcam.", dict(size=12.5, color=GREY))], first=True, sa=0, ls=1.05)
    notes(s, "Una stessa attività allena cose diverse a seconda del comando. La webcam è la parte più originale: dettagli "
              "nella prossima slide.")

# =========================================================================
# 8 — WEBCAM
# =========================================================================
def slide_webcam():
    s = header(blank(), "Comando con la webcam", "Seguire un colore, senza intelligenza artificiale")
    steps = [("1", "EVIDENZIATORE FLUO", "Il paziente tiene in mano un evidenziatore verde, giallo o fucsia"),
             ("2", "LA WEBCAM GUARDA", "Per ogni fotogramma il programma esamina i pixel dell'immagine"),
             ("3", "TROVA IL COLORE", "Tiene solo i pixel molto accesi del colore giusto: la pelle viene scartata"),
             ("4", "CALCOLA IL CENTRO", "La media di quei pixel è la posizione dell'evidenziatore"),
             ("5", "MUOVE ASTRO", "Quel punto diventa il puntatore: Astro segue il braccio (effetto specchio)")]
    x0, w, gap, y, h = 0.55, 2.30, 0.165, 1.92, 1.98
    cy = y + 0.52
    for i, (n, t, d) in enumerate(steps):
        x = x0 + i * (w + gap)
        rect(s, Inches(x), Inches(y), Inches(w), Inches(h), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
        node(s, Inches(x + w / 2 - 0.33), Inches(y + 0.16), Inches(0.66), Inches(0.66), n, fill=TEAL, ts=23, shape=MSO_SHAPE.OVAL)
        _, tf = tb(s, Inches(x + 0.12), Inches(y + 0.92), Inches(w - 0.24), Inches(0.42), anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(t, dict(size=11, bold=True, color=TEAL_DARK))], align=PP_ALIGN.CENTER, first=True, sa=0)
        _, tf2 = tb(s, Inches(x + 0.15), Inches(y + 1.32), Inches(w - 0.30), Inches(0.58))
        para(tf2, [(d, dict(size=10.5, color=INK))], align=PP_ALIGN.CENTER, first=True, sa=0, ls=1.02)
        if i < 4:
            arrow(s, Inches(x + w + 0.005), Inches(cy), Inches(x + w + gap - 0.005), Inches(cy), w=Pt(1.5))
    rect(s, Inches(0.55), Inches(4.22), Inches(6.05), Inches(2.42), GREEN_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tfa = tb(s, Inches(0.85), Inches(4.42), Inches(5.5), Inches(2.1))
    para(tfa, [("Perché un evidenziatore (e non il dito)?", dict(size=14, bold=True, color=GREEN))], first=True, sa=5)
    para(tfa, [("Riconoscere davvero una mano richiede intelligenza artificiale e librerie pesanti. Qui si insegue un "
                "colore acceso: semplice, robusto e ", dict(size=12.5, color=INK)),
               ("scritto interamente da codice, senza librerie esterne", dict(size=12.5, bold=True, color=GREEN)),
               (". I colori fluo sono molto più saturi della pelle, quindi il viso non viene confuso col bersaglio.",
                dict(size=12.5, color=INK))], sa=0, ls=1.08)
    rect(s, Inches(6.75), Inches(4.22), Inches(6.03), Inches(2.42), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tfb = tb(s, Inches(7.05), Inches(4.42), Inches(5.45), Inches(2.1))
    para(tfb, [("Regolabile per il paziente", dict(size=14, bold=True, color=TEAL_DARK))], first=True, sa=6)
    for a, b in [("Morbidezza", "calma il puntatore se la mano trema"),
                 ("Anteprima con mirino", "verde quando ti vede, rosso quando no"),
                 ("Soglie di colore/luce", "si adattano alla luce dell'ambiente"),
                 ("Mouse sempre di riserva", "il gioco non si blocca mai")]:
        para(tfb, [(a + " — ", dict(size=12, bold=True, color=INK)), (b, dict(size=12, color=GREY))],
             sa=5, ls=1.04, bullet=True, bcolor=TEAL)
    notes(s, "Punto tecnico forte: niente IA, niente librerie esterne — 'inseguo un colore'. Scelta progettuale onesta e "
              "robusta. La 'morbidezza' è il parametro pensato per il tremore.")

# =========================================================================
# 9 — FEEDBACK
# =========================================================================
def slide_feedback():
    s = header(blank(), "Il feedback al paziente", "Risposta immediata su due canali: orecchio + occhio")
    rows = [("AZIONE", "SUONO", "ASTRO (effetto visivo)"),
            ("Prende una stellina", "tre note che salgono, allegre", "si gonfia e brilla di verde"),
            ("Raggiunge la porta / vince", "piccola fanfara", "si gonfia e brilla"),
            ("Tocca asteroide o bomba", "due note basse che scendono", "si schiaccia e lampeggia di rosso"),
            ("Tempo scaduto", "come l'errore (gentile)", "si schiaccia")]
    tx, ty = 0.55, 1.95
    rh = 0.76; cw = [2.75, 2.55, 2.45]
    yy = ty
    for ri, row in enumerate(rows):
        head = ri == 0; xx = tx
        for ci, cell in enumerate(row):
            if head:
                fill = TEAL_DARK
            else:
                fill = (GREEN_LT if ri in (1, 2) else RED_LT) if ci > 0 else WHITE
            cellsh = rect(s, Inches(xx), Inches(yy), Inches(cw[ci]), Inches(rh), fill, line=WHITE, line_w=Pt(1.5))
            tf = cellsh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.10); tf.margin_right = Inches(0.05)
            para(tf, [(cell, dict(size=12 if head else 11.5, bold=head or ci == 0, color=WHITE if head else INK))],
                 first=True, sa=0, ls=1.0)
            xx += cw[ci]
        yy += rh
    rect(s, Inches(8.55), Inches(1.95), Inches(4.23), Inches(3.78), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tf = tb(s, Inches(8.85), Inches(2.18), Inches(3.65), Inches(3.4))
    para(tf, [("Perché serve in terapia", dict(size=15, bold=True, color=TEAL_DARK))], first=True, sa=7)
    para(tf, [("Un riscontro immediato e su due sensi insieme aiuta il paziente a capire subito se il movimento è "
               "corretto, rinforzando i gesti giusti.", dict(size=12.5, color=INK))], sa=8, ls=1.08)
    para(tf, [("Il suono dell'errore è apposta più basso e gentile di quello positivo: deve ", dict(size=12.5, color=INK)),
              ("informare, non spaventare.", dict(size=12.5, bold=True, color=ORANGE))], sa=8, ls=1.08)
    para(tf, [("Suono ed effetti si possono spegnere per pazienti sensibili o ambienti silenziosi.",
               dict(size=12.5, color=INK))], sa=0, ls=1.08)
    notes(s, "Concetto di apprendimento motorio: il 'feedback aumentato' immediato consolida il gesto corretto. Doppio "
              "canale (audio+video) per chi ha un canale più debole. Errore = segnale gentile, non punitivo.")

# =========================================================================
# 10 — LIVELLI + PARAMETRI
# =========================================================================
def slide_livelli():
    s = header(blank(), "Livelli e parametri", "Difficoltà crescente, tutto regolabile")
    livs = [("LIVELLO 1", "Primo volo", "Spazio aperto, nessun ostacolo: solo stelline da raccogliere. Si prende confidenza col comando.", TEAL),
            ("LIVELLO 2", "Tra gli asteroidi", "Compaiono gli asteroidi (muri che fanno male): servono traiettorie più precise.", TEAL_DARK),
            ("LIVELLO 3", "Campo minato", "Asteroidi più quattro bombe da evitare: massima attenzione e controllo.", AMBER)]
    x0, w, gap = 0.55, 3.95, 0.19
    for i, (t, st, d, col) in enumerate(livs):
        x = x0 + i * (w + gap)
        rect(s, Inches(x), Inches(1.82), Inches(w), Inches(2.30), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=GREY_LT)
        rect(s, Inches(x), Inches(1.82), Inches(0.16), Inches(2.30), (ORANGE if col == AMBER else col))
        _, tf = tb(s, Inches(x + 0.32), Inches(2.02), Inches(w - 0.5), Inches(2.0))
        para(tf, [(t, dict(size=12, bold=True, color=(ORANGE if col == AMBER else col)))], first=True, sa=1)
        para(tf, [(st, dict(size=18, bold=True, color=INK))], sa=7)
        para(tf, [(d, dict(size=12.5, color=GREY))], sa=0, ls=1.08)
    rect(s, Inches(0.55), Inches(4.40), Inches(12.23), Inches(2.22), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
    _, tfh = tb(s, Inches(0.85), Inches(4.56), Inches(11.6), Inches(0.5))
    para(tfh, [("La «pagina dei valori» — ", dict(size=14, bold=True, color=TEAL_DARK)),
               ("difficoltà e accessibilità si cambiano modificando dei numeri, senza toccare la logica del gioco:",
                dict(size=13.5, color=INK))], first=True, sa=0)
    chips = [("Vite per livello", "5"), ("Tempo per livello", "60 s"), ("Stelline (liv. 1)", "10"), ("Punti per stellina", "10"),
             ("Morbidezza webcam", "anti-tremore"), ("Volume errore", "più basso"), ("Suono / effetti", "on / off"), ("Pulsanti dwell", "1,2 s")]
    cw2, ch = 2.86, 0.58
    for i, (k, v) in enumerate(chips):
        cx = 0.85 + (i % 4) * 2.98
        cyy = 5.22 + (i // 4) * 0.68
        chip(s, Inches(cx), Inches(cyy), Inches(cw2), Inches(ch), k, v)
    notes(s, "Tre livelli a difficoltà crescente. Tutti i numeri (difficoltà, tempi, volumi, anti-tremore) vivono in poche "
              "'pagine dei parametri': il terapista può tarare il gioco sul singolo paziente senza programmare.")

# =========================================================================
# 11 — COSA ABBIAMO FATTO
# =========================================================================
def slide_lavoro():
    s = header(blank(), "Il lavoro svolto", "Cosa ho realizzato (in sintesi, senza codice)")
    cols = [
        ("IL GIOCO", TEAL, [
            "Motore di gioco 2D completo in Unity / C#",
            "Grafica e suoni generati interamente da codice",
            "Tre livelli a difficoltà crescente con stelline, asteroidi e bombe",
            "Schermata iniziale e pannello con i parametri tecnici del PC"]),
        ("L'ACCESSIBILITÀ", AMBER, [
            "Tre modalità di comando: mouse, webcam, joystick",
            "Tracciamento del colore via webcam (senza IA né librerie)",
            "Pulsanti 'a permanenza' (dwell): si usano senza clic",
            "Si torna da soli al mouse se la webcam non c'è"]),
        ("LA TERAPIA", GREEN, [
            "Feedback immediato sonoro e visivo per ogni azione",
            "Parametri di difficoltà e accessibilità centralizzati",
            "Anti-tremore (morbidezza) e volumi regolabili",
            "Compatibilità Windows / Mac per l'uso reale"]),
    ]
    x0, w, gap = 0.55, 3.97, 0.18
    for i, (t, col, items) in enumerate(cols):
        x = x0 + i * (w + gap)
        rect(s, Inches(x), Inches(1.85), Inches(w), Inches(4.85), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=GREY_LT)
        rect(s, Inches(x), Inches(1.85), Inches(w), Inches(0.62), col, shape=MSO_SHAPE.ROUND_2_SAME_RECTANGLE)
        _, tf = tb(s, Inches(x + 0.18), Inches(1.94), Inches(w - 0.36), Inches(0.46), anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(t, dict(size=14, bold=True, color=(AMBER_INK if col == AMBER else WHITE)))], align=PP_ALIGN.CENTER, first=True, sa=0)
        _, tf2 = tb(s, Inches(x + 0.24), Inches(2.75), Inches(w - 0.48), Inches(3.8))
        first = True
        for it in items:
            para(tf2, [(it, dict(size=12.5, color=INK))], first=first, sa=9, bullet=True,
                 bcolor=(ORANGE if col == AMBER else col), ls=1.05); first = False
    notes(s, "Riepilogo del lavoro, senza codice. Tre filoni: il gioco, l'accessibilità (la parte più originale) e gli "
              "agganci alla terapia. Tutto generato da codice = nessuna dipendenza esterna, massima portabilità.")

# =========================================================================
# 12 — RIABILITAZIONE (LIT)
# =========================================================================
def slide_riabilitazione():
    s = header(blank(), "La riabilitazione", "Di che terapia si tratta")
    rect(s, Inches(0.55), Inches(1.82), Inches(6.55), Inches(1.95), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tf = tb(s, Inches(0.85), Inches(2.02), Inches(5.95), Inches(1.6))
    para(tf, [("Tipo di riabilitazione", dict(size=15, bold=True, color=TEAL_DARK))], first=True, sa=6)
    para(tf, [(LIT.get("tipoRiabilitazione", ""), dict(size=13, color=INK))], sa=0, ls=1.1)
    # come il gioco la supporta
    _, tf2 = tb(s, Inches(0.55), Inches(4.05), Inches(6.55), Inches(2.7))
    para(tf2, [("Come il gioco la supporta", dict(size=15, bold=True, color=TEAL_DARK))], first=True, sa=8)
    supp = [("Gesti ampi del braccio", "la modalità webcam sposta l'esercizio dal polso a tutto l'arto superiore"),
            ("Ripetizione senza noia", "raccogliere stelline = tante ripetizioni del gesto, ma motivate"),
            ("Coordinazione occhio-mano", "guidare Astro verso bersagli e porta allena mira e controllo"),
            ("Dosaggio su misura", "difficoltà, tempi e anti-tremore regolabili per ogni paziente")]
    for a, b in supp:
        para(tf2, [(a + " — ", dict(size=12.5, bold=True, color=INK)), (b, dict(size=12.5, color=GREY))],
             sa=7, bullet=True, bcolor=TEAL, ls=1.06)
    # destra: risultati attesi
    rect(s, Inches(7.40), Inches(1.82), Inches(5.38), Inches(4.9), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
    _, tfr = tb(s, Inches(7.70), Inches(2.05), Inches(4.85), Inches(0.5))
    para(tfr, [("Cosa ci si aspetta", dict(size=15, bold=True, color=TEAL_DARK))], first=True, sa=0)
    yy = 2.62
    for r in LIT.get("risultatiAttesi", [])[:6]:
        node(s, Inches(7.72), Inches(yy + 0.02), Inches(0.28), Inches(0.28), "✓", fill=GREEN, ts=12, shape=MSO_SHAPE.OVAL)
        _, tfx = tb(s, Inches(8.14), Inches(yy - 0.06), Inches(4.4), Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
        para(tfx, [(r, dict(size=12.5, color=INK))], first=True, sa=0, ls=1.05)
        yy += max(0.62, 0.42 + 0.20 * (len(r) // 38))
    _, tfn = tb(s, Inches(7.70), Inches(6.18), Inches(4.9), Inches(0.5))
    para(tfn, [("Nota: prototipo a scopo didattico; l'efficacia clinica va validata con uno studio dedicato.",
                dict(size=10.5, italic=True, color=GREY))], first=True, sa=0, ls=1.0)
    notes(s, "Inquadra il tipo di terapia (riabilitazione motoria dell'arto superiore / coordinazione) e collega ogni "
              "caratteristica del gioco a un obiettivo terapeutico. Sii onesto: è un prototipo, l'efficacia va validata.")

# =========================================================================
# 13 — VANTAGGI + EVIDENZE (LIT)
# =========================================================================
def slide_evidenze():
    s = header(blank(), "Perché i giochi aiutano", "Vantaggi ed evidenze dalla letteratura")
    # sinistra: vantaggi (icone)
    _, tf = tb(s, Inches(0.55), Inches(1.80), Inches(5.7), Inches(0.4))
    para(tf, [("I vantaggi dei serious game in età pediatrica", dict(size=14.5, bold=True, color=TEAL_DARK))], first=True, sa=0)
    yy = 2.35
    for v in LIT.get("vantaggi", [])[:6]:
        node(s, Inches(0.6), Inches(yy + 0.02), Inches(0.30), Inches(0.30), "★", fill=AMBER, ts=11, shape=MSO_SHAPE.OVAL)
        _, tfx = tb(s, Inches(1.05), Inches(yy - 0.07), Inches(5.2), Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
        # v può essere "Titolo — spiegazione"
        if "—" in v:
            a, b = v.split("—", 1)
            para(tfx, [(a.strip() + " — ", dict(size=12.5, bold=True, color=INK)), (b.strip(), dict(size=12.5, color=GREY))],
                 first=True, sa=0, ls=1.05)
        else:
            para(tfx, [(v, dict(size=12.5, color=INK))], first=True, sa=0, ls=1.05)
        yy += max(0.66, 0.44 + 0.22 * (len(v) // 46))
    # destra: evidenze con riferimento
    rect(s, Inches(6.55), Inches(1.80), Inches(6.23), Inches(4.95), PANEL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
    _, tfe = tb(s, Inches(6.85), Inches(2.02), Inches(5.7), Inches(0.4))
    para(tfe, [("Cosa dice la ricerca", dict(size=14.5, bold=True, color=TEAL_DARK))], first=True, sa=0)
    _, tfb = tb(s, Inches(6.85), Inches(2.52), Inches(5.7), Inches(4.1))
    first = True
    for e in LIT.get("evidenze", [])[:5]:
        para(tfb, [(e.get("punto", ""), dict(size=12, color=INK))], first=first, sa=1, bullet=True, bcolor=TEAL, ls=1.05)
        para(tfb, [("       " + e.get("riferimento", ""), dict(size=10.5, italic=True, color=TEAL_DARK))], sa=8, ls=1.0)
        first = False
    notes(s, "Qui porti le prove: i serious game/exergame aumentano motivazione e quantità di pratica, e il feedback "
              "immediato favorisce l'apprendimento motorio. Cita gli autori mostrati. La bibliografia completa è nella tesina.")

# =========================================================================
# 14 — STRUTTURA DELLA TESINA
# =========================================================================
def slide_tesina():
    s = header(blank(), "La tesina", "Come strutturare la relazione scritta")
    _, tfi = tb(s, Inches(0.55), Inches(1.78), Inches(12.2), Inches(0.5))
    para(tfi, [("La tesina approfondisce ciò che le slide mostrano in sintesi: testo, immagini del gioco e riferimenti "
                "scientifici. Struttura consigliata:", dict(size=13.5, color=INK))], first=True, sa=0, ls=1.06)
    secs = [("1 · Introduzione", "il problema della motivazione in riabilitazione e l'idea del gioco"),
            ("2 · Il gioco «Astro»", "concept, livelli, regole — con schermate"),
            ("3 · Architettura software", "i moduli e le scelte di progetto (schema a blocchi)"),
            ("4 · Modalità di comando", "mouse, webcam (tracking del colore), joystick"),
            ("5 · Feedback al paziente", "risposta audiovisiva e apprendimento motorio"),
            ("6 · La riabilitazione", "tipo di terapia e popolazione di riferimento"),
            ("7 · Vantaggi ed evidenze", "rassegna della letteratura sui serious game"),
            ("8 · Risultati attesi e metriche", "cosa misurare per valutarne l'utilità"),
            ("9 · Conclusioni e sviluppi", "limiti e prossimi passi"),
            ("Bibliografia + Appendice", "riferimenti citati e pagina dei parametri")]
    for i, (t, d) in enumerate(secs):
        col = i // 5; row = i % 5
        x = 0.55 + col * 6.25; y = 2.45 + row * 0.86
        rect(s, Inches(x), Inches(y), Inches(6.0), Inches(0.74), WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=TEAL_LT2)
        _, tf = tb(s, Inches(x + 0.18), Inches(y), Inches(5.7), Inches(0.74), anchor=MSO_ANCHOR.MIDDLE)
        para(tf, [(t + "  ", dict(size=12.5, bold=True, color=TEAL_DARK)), ("— " + d, dict(size=11.5, color=GREY))],
             first=True, sa=0, ls=1.0)
    notes(s, "Questa slide spiega al prof come è organizzata la relazione scritta e mostra il legame PP↔tesina. "
              "La tesina in Word (file a parte) contiene già questi capitoli con segnaposto per le immagini.")

# =========================================================================
# 15 — CONCLUSIONI
# =========================================================================
def slide_conclusioni():
    s = header(blank(), "Conclusioni", "Cosa funziona e dove si può arrivare")
    rect(s, Inches(0.55), Inches(1.85), Inches(6.05), Inches(4.85), GREEN_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tf = tb(s, Inches(0.85), Inches(2.08), Inches(5.5), Inches(4.4))
    para(tf, [("Cosa ho ottenuto", dict(size=16, bold=True, color=GREEN))], first=True, sa=8)
    for it in ["Un'interfaccia software completa e funzionante, scritta da zero",
               "Tre modi di comandare il gioco, incluso il movimento del braccio via webcam",
               "Feedback immediato e parametri regolabili pensati per la terapia",
               "Nessuna dipendenza esterna: gira su Windows e Mac"]:
        para(tf, [(it, dict(size=13, color=INK))], sa=9, bullet=True, bcolor=GREEN, ls=1.07)
    rect(s, Inches(6.95), Inches(1.85), Inches(5.83), Inches(4.85), TEAL_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    _, tf2 = tb(s, Inches(7.25), Inches(2.08), Inches(5.3), Inches(4.4))
    para(tf2, [("Sviluppi futuri", dict(size=16, bold=True, color=TEAL_DARK))], first=True, sa=8)
    for it in ["Salvare i progressi e misurare metriche (tempi, errori, ampiezza dei movimenti)",
               "Più livelli ed esercizi mirati a gesti specifici",
               "Calibrazione guidata della webcam per ogni paziente",
               "Validazione con terapisti e piccoli studi pilota",
               "Profili paziente e report per il terapista"]:
        para(tf2, [(it, dict(size=13, color=INK))], sa=9, bullet=True, bcolor=TEAL, ls=1.07)
    notes(s, "Chiudi con onestà: prototipo solido dal lato ingegneristico; il passo successivo è misurare e validare. "
              "Le metriche (tempi, errori, ampiezza dei movimenti) sono il ponte verso un uso clinico.")

# =========================================================================
# 16 — CHIUSURA
# =========================================================================
def slide_chiusura():
    s = blank()
    rect(s, 0, 0, SW, SH, TEAL_DEEP)
    rect(s, 0, Inches(7.34), SW, Inches(0.16), AMBER)
    for cx, cy, d, col in [(11.6, 1.0, 0.7, TEAL), (10.8, 6.0, 0.4, TEAL_MID), (1.4, 5.9, 0.5, TEAL),
                           (2.0, 1.2, 0.30, AMBER), (12.4, 4.0, 0.24, WHITE)]:
        rect(s, Inches(cx), Inches(cy), Inches(d), Inches(d), col, shape=MSO_SHAPE.OVAL)
    node(s, Inches(5.86), Inches(1.75), Inches(1.6), Inches(1.6), "★", fill=AMBER, ts=58, shape=MSO_SHAPE.OVAL)
    _, tf = tb(s, Inches(1.5), Inches(3.65), Inches(10.33), Inches(2.0))
    para(tf, [("Grazie per l'attenzione", dict(size=36, bold=True, color=WHITE))], align=PP_ALIGN.CENTER, first=True, sa=6)
    para(tf, [("Astro — un videogioco per la riabilitazione motoria", dict(size=18, color=TEAL_LT))],
         align=PP_ALIGN.CENTER, sa=4)
    para(tf, [("Domande?", dict(size=16, italic=True, color=AMBER))], align=PP_ALIGN.CENTER, sb=8)
    _, tfp = tb(s, Inches(1.5), Inches(6.55), Inches(10.33), Inches(0.5))
    para(tfp, [("Politecnico di Bari", dict(size=13, bold=True, color=WHITE))], align=PP_ALIGN.CENTER, first=True, sa=0)
    notes(s, "Ringrazia e apri alle domande. Tieni pronte: 'perché senza IA?', 'come misureresti l'efficacia?', "
              "'che ruolo ha il terapista?'.")

# ---------- Build ----------
for fn in [slide_cover, slide_indice, slide_obiettivo, slide_concept, slide_flusso, slide_architettura,
           slide_comandi, slide_webcam, slide_feedback, slide_livelli, slide_lavoro, slide_riabilitazione,
           slide_evidenze, slide_tesina, slide_conclusioni, slide_chiusura]:
    fn()

prs.save(OUT)

# ---------- Controllo geometria: shape fuori slide ----------
def emu(v): return int(v)
warn = 0
for idx, sl in enumerate(prs.slides, 1):
    for sh in sl.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
            if l is None: continue
            if l < -9000 or t < -9000 or (l + (w or 0)) > SW + 9000 or (t + (h or 0)) > SH + 9000:
                nm = getattr(sh, "name", "?")
                print(f"  ! slide {idx}: '{nm}' fuori bordo  L={l/914400:.2f} T={t/914400:.2f} R={(l+(w or 0))/914400:.2f} B={(t+(h or 0))/914400:.2f}")
                warn += 1
        except Exception:
            pass
print(f"OK: {len(prs.slides.__iter__.__self__._sldIdLst)} slide salvate in {OUT}  (avvisi geometria: {warn})")
