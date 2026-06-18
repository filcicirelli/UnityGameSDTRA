# -*- coding: utf-8 -*-
"""Genera 'Tesina_Astro.docx' — relazione scritta del progetto (Politecnico di Bari)."""
import os, json
HERE = os.path.dirname(os.path.abspath(__file__))
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = os.environ.get("OUT_DOCX", os.path.join(HERE, "Tesina_Astro.docx"))
LOGO = os.path.join(HERE, "logo_polibari.png")
LIT_PATH = os.path.join(HERE, "lit_docx.json")

if os.path.exists(LIT_PATH):
    with open(LIT_PATH, encoding="utf-8") as f:
        LIT = json.load(f)
else:
    LIT = {"tipoRiabilitazione": "(placeholder)", "vantaggi": [], "evidenze": [],
           "risultatiAttesi": [], "bibliografia": []}

TEAL_DEEP = RGBColor(0x00, 0x4C, 0x4B)
TEAL_DARK = RGBColor(0x00, 0x6E, 0x6C)
GREY = RGBColor(0x5E, 0x6E, 0x6E)
INK = RGBColor(0x21, 0x2B, 0x2B)

doc = Document()

# ---- Stili base ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11.5)
normal.font.color.rgb = INK
pf = normal.paragraph_format
pf.space_after = Pt(7); pf.line_spacing = 1.18

for lvl, sz, col in [("Title", 30, TEAL_DEEP), ("Heading 1", 18, TEAL_DEEP),
                     ("Heading 2", 14, TEAL_DARK), ("Heading 3", 12, TEAL_DARK)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"; st.font.size = Pt(sz); st.font.bold = True
    st.font.color.rgb = col

for sec in doc.sections:
    sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)

def p(text="", style=None, size=None, bold=False, italic=False, color=None, align=None, after=None, before=None):
    par = doc.add_paragraph(style=style)
    if align is not None: par.alignment = align
    if after is not None: par.paragraph_format.space_after = Pt(after)
    if before is not None: par.paragraph_format.space_before = Pt(before)
    if text:
        r = par.add_run(text)
        if size: r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic
        if color: r.font.color.rgb = color
    return par

def runs(par, parts):
    for t, o in parts:
        r = par.add_run(t)
        if "size" in o: r.font.size = Pt(o["size"])
        r.font.bold = o.get("bold", False); r.font.italic = o.get("italic", False)
        if "color" in o: r.font.color.rgb = o["color"]
    return par

def bullet(text_or_parts, bold_lead=None):
    par = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        runs(par, [(bold_lead, dict(bold=True, color=INK)), (text_or_parts, dict(color=INK))])
    else:
        par.add_run(text_or_parts)
    return par

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)

def image_placeholder(caption, h_cm=6.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Cm(15)
    shade(cell, "F2F7F6")
    # bordo
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "dashed"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "8FBFBC")
        borders.append(e)
    tcPr.append(borders)
    cp = cell.paragraphs[0]; cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(int(h_cm * 9)); cp.paragraph_format.space_after = Pt(int(h_cm * 9))
    r = cp.add_run("🖼  " + caption); r.font.italic = True; r.font.size = Pt(10.5); r.font.color.rgb = GREY
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rc = cap.add_run(caption); rc.font.italic = True; rc.font.size = Pt(9.5); rc.font.color.rgb = GREY
    cap.paragraph_format.space_after = Pt(10)

def add_toc():
    par = doc.add_paragraph()
    run = par.add_run()
    fldBegin = OxmlElement("w:fldChar"); fldBegin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldSep = OxmlElement("w:fldChar"); fldSep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = "Aggiorna i campi (F9) per generare l'indice."
    fldEnd = OxmlElement("w:fldChar"); fldEnd.set(qn("w:fldCharType"), "end")
    for el in (fldBegin, instr, fldSep, t, fldEnd):
        run._r.append(el)

# ============================ FRONTESPIZIO ============================
try:
    pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(LOGO, width=Cm(7))
except Exception:
    p("POLITECNICO DI BARI", bold=True, size=18, color=TEAL_DEEP, align=WD_ALIGN_PARAGRAPH.CENTER)
p("Politecnico di Bari", size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, before=4)
p("Corso di Programmazione / Informatica", size=12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)

p("ASTRO", size=34, bold=True, color=TEAL_DEEP, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
p("Un videogioco per la riabilitazione motoria", size=18, bold=True, color=TEAL_DARK,
  align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
p("Progettazione e sviluppo di un'interfaccia software che trasforma l'esercizio riabilitativo in gioco, "
  "con comando del personaggio anche tramite i movimenti del braccio davanti alla webcam.",
  size=12, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=46)

for k, v in [("Studente", "______________________________"),
             ("Relatore", "______________________________"),
             ("Anno Accademico", "20__ / 20__")]:
    par = p(align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    runs(par, [(k + ":  ", dict(bold=True, color=TEAL_DARK, size=12)), (v, dict(color=INK, size=12))])

doc.add_page_break()

# ============================ SOMMARIO ============================
doc.add_heading("Sommario", level=1)
p("Questo lavoro presenta «Astro», un videogioco bidimensionale progettato e sviluppato da zero come interfaccia "
  "di supporto alla riabilitazione motoria. Il paziente guida un piccolo alieno per raccogliere stelline ed evitare "
  "ostacoli; il personaggio può essere comandato con il mouse, con un joystick oppure muovendo un evidenziatore fluo "
  "davanti alla webcam, così da allenare gesti ampi e controllati dell'arto superiore. Ogni azione riceve un feedback "
  "immediato sonoro e visivo, e tutti i parametri di difficoltà e accessibilità sono raccolti in pagine dedicate, "
  "regolabili senza modificare la logica del programma. La relazione descrive il gioco, l'architettura del software, "
  "le scelte progettuali e il razionale riabilitativo, con una rassegna della letteratura sui serious game in età "
  "pediatrica e una proposta di metriche per valutarne l'utilità.")

p("Parole chiave: ", bold=True, after=2)
par = doc.paragraphs[-1]
par.add_run("riabilitazione motoria, serious game, exergame, tracciamento del colore, webcam, apprendimento motorio, "
            "feedback aumentato, accessibilità.").italic = True

doc.add_heading("Indice", level=1)
add_toc()
doc.add_page_break()

# ============================ 1. INTRODUZIONE ============================
doc.add_heading("1. Introduzione", level=1)
doc.add_heading("1.1 Il contesto e il problema", level=2)
p("La riabilitazione motoria si basa sulla ripetizione di esercizi: per riacquistare o migliorare il controllo di un "
  "movimento occorre ripeterlo molte volte, in modo corretto e via via più impegnativo. Questa ripetizione, però, è "
  "spesso percepita come monotona, e l'aderenza al trattamento — cioè la costanza con cui il paziente svolge gli "
  "esercizi — tende a calare. Il problema è particolarmente sentito con i bambini, per i quali la motivazione è un "
  "fattore decisivo del risultato.")
doc.add_heading("1.2 L'idea", level=2)
p("L'idea alla base del progetto è semplice: trasformare l'esercizio in un gioco. Se il movimento da allenare diventa "
  "il modo per raccogliere stelline ed evitare asteroidi, il paziente continua a muoversi — e quindi ad allenarsi — "
  "mentre si diverte. Il movimento può essere eseguito con il braccio nello spazio, inquadrato da una comune webcam, "
  "allenando gesti ampi e controllati anziché i piccoli movimenti del polso richiesti dal mouse.")
doc.add_heading("1.3 Obiettivo del lavoro", level=2)
p("L'obiettivo è duplice: da un lato dimostrare competenze concrete di programmazione, progettando e realizzando da "
  "zero un'applicazione interattiva completa; dall'altro, far sì che questa applicazione sia realmente sensata in un "
  "contesto riabilitativo, con scelte pensate per l'accessibilità e per il terapista. Il presente prototipo non è uno "
  "strumento clinico validato, ma una base solida su cui costruire e da sottoporre, in seguito, a verifica sperimentale.")

# ============================ 2. IL GIOCO ============================
doc.add_heading("2. Il gioco «Astro»", level=1)
doc.add_heading("2.1 Concept", level=2)
p("Il giocatore controlla «Astro», un alieno a bordo di un piccolo UFO. Astro segue il puntatore: dove punta il "
  "giocatore, lì si sposta. Ogni livello si svolge in due fasi: nella prima si devono raccogliere tutte le stelline "
  "presenti nello spazio di gioco; quando sono state raccolte, compare una porta e basta raggiungerla per completare "
  "il livello. Durante il percorso si devono evitare gli ostacoli, che fanno perdere una vita.")
image_placeholder("Figura 2.1 — Schermata di gioco: Astro, le stelline e lo sfondo spaziale.")
doc.add_heading("2.2 Regole e livelli", level=2)
p("Il giocatore dispone di un numero fisso di vite e di un tempo limite per ogni livello. Toccare un asteroide o una "
  "bomba, oppure lasciare scadere il tempo, costa una vita; esaurite le vite si ha un «game over» con possibilità di "
  "riprovare. I tre livelli hanno difficoltà crescente:")
bullet("spazio aperto, nessun ostacolo: si prende confidenza con il comando.", bold_lead="Livello 1 — Primo volo: ")
bullet("compaiono gli asteroidi (muri che fanno male), che impongono traiettorie più precise.",
       bold_lead="Livello 2 — Tra gli asteroidi: ")
bullet("agli asteroidi si aggiungono alcune bombe da evitare: massima attenzione e controllo.",
       bold_lead="Livello 3 — Campo minato: ")
p("Una breve fase iniziale «pronti…» e una brevissima invulnerabilità dopo ogni colpo evitano che il paziente perda "
  "tutte le vite per un singolo errore involontario, scelta importante per chi ha un controllo motorio ancora incerto.")
image_placeholder("Figura 2.2 — I tre livelli a difficoltà crescente.")

# ============================ 3. ARCHITETTURA ============================
doc.add_heading("3. L'architettura del software", level=1)
p("Il programma è sviluppato in Unity con linguaggio C#. Una scelta progettuale trasversale è che tutto — grafica e "
  "suoni — è generato da codice: non esistono file grafici o audio esterni. Questo rende l'applicazione leggera, "
  "facilmente portabile e completamente sotto controllo, ed è anche una dimostrazione di competenza tecnica.")
doc.add_heading("3.1 I moduli principali", level=2)
bullet("prepara la scena all'avvio (telecamera) e crea i moduli principali.", bold_lead="Avvio: ")
bullet("il «cervello» del gioco: tiene stato, punteggio, vite, tempo e livello, e decide cosa accade a ogni evento.",
       bold_lead="Gestore di gioco: ")
bullet("punto unico che risponde alla domanda «dove punta il paziente?», a seconda della modalità scelta.",
       bold_lead="Comandi: ")
bullet("il personaggio del giocatore; segue il puntatore e rileva il contatto con stelline, ostacoli e porta.",
       bold_lead="Astro: ")
bullet("costruiscono i tre livelli e tutti gli oggetti in scena.", bold_lead="Caricatore e Definizione livelli: ")
bullet("produce la risposta sonora e visiva alle azioni del paziente.", bold_lead="Feedback paziente: ")
bullet("disegna l'interfaccia a schermo (HUD), gestisce il cursore e i pulsanti accessibili.",
       bold_lead="Interfaccia di gioco: ")
bullet("genera via codice tutte le immagini del gioco.", bold_lead="Fabbrica immagini: ")
image_placeholder("Figura 3.1 — Schema a blocchi dell'architettura (inserire la slide corrispondente).", h_cm=7)
doc.add_heading("3.2 Le scelte di progetto", level=2)
p("Il principio guida è la separazione delle responsabilità: ogni modulo fa una cosa sola. Il modulo «Comandi» è "
  "emblematico: Astro non sa quale dispositivo lo stia guidando, chiede soltanto «dove devo andare?». Per questo "
  "aggiungere una nuova modalità di comando non richiede di modificare il resto del gioco. Allo stesso modo, tutti i "
  "valori regolabili (difficoltà, tempi, volumi, sensibilità) sono raccolti in apposite pagine di parametri, separate "
  "dalla logica: il comportamento del gioco si adatta cambiando dei numeri, non riscrivendo il codice.")

# ============================ 4. COMANDI ============================
doc.add_heading("4. Le modalità di comando", level=1)
p("Il gioco offre tre modi di comandare Astro, così da adattarsi alle capacità del paziente e all'attrezzatura "
  "disponibile. La modalità si sceglie nella schermata iniziale.")
bullet("controllo fine del polso e della mano; sempre disponibile come riserva.", bold_lead="Mouse: ")
bullet("si muove un evidenziatore fluo davanti alla webcam per allenare gesti ampi del braccio nello spazio.",
       bold_lead="Dito (webcam): ")
bullet("comando fisico con leva, frecce o tasti, utile per chi usa già ausili.", bold_lead="Joystick: ")
doc.add_heading("4.1 Il comando tramite webcam (tracciamento del colore)", level=2)
p("La modalità più originale non utilizza intelligenza artificiale né librerie esterne. Il paziente tiene in mano un "
  "evidenziatore fluo (verde, giallo o fucsia); per ogni fotogramma il programma esamina i pixel dell'immagine, "
  "mantiene solo quelli del colore giusto e molto accesi (la pelle e il viso, meno saturi, vengono scartati), ne "
  "calcola il centro e lo trasforma nella posizione del puntatore, con effetto specchio. Da quel momento Astro segue "
  "il movimento del braccio esattamente come seguirebbe il mouse.")
p("La scelta di inseguire un colore acceso, invece di riconoscere una mano, è deliberata: è semplice, robusta e "
  "interamente realizzabile in codice. La modalità è inoltre regolabile per il paziente: un parametro di «morbidezza» "
  "calma il puntatore quando la mano trema, un'anteprima con mirino mostra in tempo reale se l'evidenziatore è "
  "tracciato correttamente, e se la webcam manca il gioco torna automaticamente al mouse senza bloccarsi.")
image_placeholder("Figura 4.1 — Modalità webcam: anteprima con mirino e tracciamento dell'evidenziatore.")
doc.add_heading("4.2 Pulsanti accessibili (dwell)", level=2)
p("Senza mouse non c'è il clic: i pulsanti dei pannelli (per esempio «prossimo livello» o «riprova») si attivano "
  "tenendo Astro sopra il pulsante per un breve tempo, durante il quale una barra si riempie. Questo evita di restare "
  "bloccati nei passaggi tra una schermata e l'altra quando si gioca con la webcam o il joystick.")

# ============================ 5. FEEDBACK ============================
doc.add_heading("5. Il feedback al paziente", level=1)
p("A ogni azione il gioco risponde su due canali sensoriali insieme — udito e vista — in modo immediato. La tabella "
  "riassume le risposte:")
tbl = doc.add_table(rows=5, cols=3); tbl.style = "Light Grid Accent 1"; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
data = [("Azione", "Suono", "Reazione di Astro"),
        ("Prende una stellina", "tre note che salgono, allegre", "si gonfia e brilla di verde"),
        ("Raggiunge la porta / vince", "piccola fanfara", "si gonfia e brilla"),
        ("Tocca asteroide o bomba", "due note basse che scendono", "si schiaccia e lampeggia di rosso"),
        ("Tempo scaduto", "come l'errore, ma gentile", "si schiaccia")]
for i, row in enumerate(data):
    for j, val in enumerate(row):
        c = tbl.cell(i, j); c.text = ""
        rp = c.paragraphs[0]; rr = rp.add_run(val)
        rr.font.size = Pt(10.5)
        if i == 0:
            rr.font.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "006E6C")
p("")
p("Il razionale è quello del feedback aumentato dell'apprendimento motorio: un'informazione di ritorno immediata e "
  "ridondante su più canali aiuta il paziente a capire subito se il movimento è corretto e a consolidare il gesto "
  "giusto. Il suono dell'errore è volutamente più basso e gentile di quello positivo: deve informare, non spaventare. "
  "Suono ed effetti visivi possono essere disattivati per pazienti sensibili o ambienti silenziosi.")

# ============================ 6. RIABILITAZIONE ============================
doc.add_heading("6. La riabilitazione di riferimento", level=1)
doc.add_heading("6.1 Tipo di riabilitazione", level=2)
p(LIT.get("tipoRiabilitazione", ""))
doc.add_heading("6.2 Come il gioco supporta la terapia", level=2)
for a, b in [("Gesti ampi dell'arto superiore", "la modalità webcam sposta l'esercizio dal polso a tutto il braccio."),
             ("Ripetizione motivata", "raccogliere le stelline equivale a ripetere molte volte il gesto, ma con uno scopo di gioco."),
             ("Coordinazione occhio-mano", "guidare Astro verso bersagli e porta allena mira, controllo e pianificazione del movimento."),
             ("Dosaggio individuale", "difficoltà, tempi, volumi e sensibilità si regolano per ciascun paziente.")]:
    bullet(b, bold_lead=a + ": ")

# ============================ 7. EVIDENZE ============================
doc.add_heading("7. I serious game in riabilitazione: vantaggi ed evidenze", level=1)
p("Negli ultimi anni i cosiddetti serious game ed exergame — videogiochi con finalità non solo di intrattenimento ma "
  "anche terapeutiche o educative — sono stati studiati come strumento di supporto alla riabilitazione, in particolare "
  "in età pediatrica. Di seguito i principali vantaggi e alcune evidenze dalla letteratura.")
doc.add_heading("7.1 Vantaggi", level=2)
for v in LIT.get("vantaggi", []):
    if "—" in v:
        a, b = v.split("—", 1); bullet(b.strip(), bold_lead=a.strip() + " — ")
    else:
        bullet(v)
doc.add_heading("7.2 Cosa dice la ricerca", level=2)
for e in LIT.get("evidenze", []):
    par = bullet(e.get("punto", ""))
    if e.get("riferimento"):
        rr = par.add_run("  [" + e["riferimento"] + "]"); rr.font.italic = True; rr.font.color.rgb = TEAL_DARK; rr.font.size = Pt(10.5)
p("Le evidenze vanno lette con cautela: la qualità degli studi è eterogenea e i risultati dipendono molto dal tipo di "
  "gioco, dalla dose di trattamento e dalla popolazione. Tuttavia, l'indicazione ricorrente è che un gioco ben "
  "progettato aumenta la motivazione e la quantità di pratica e che il feedback immediato favorisce l'apprendimento "
  "motorio — esattamente i meccanismi su cui «Astro» fa leva.")

# ============================ 8. RISULTATI ATTESI ============================
doc.add_heading("8. Risultati attesi e metriche di valutazione", level=1)
p("Per passare da prototipo a strumento utile occorre misurare. Tra i risultati attesi e gli indicatori che si "
  "potrebbero raccogliere automaticamente durante il gioco:")
for r in LIT.get("risultatiAttesi", []):
    bullet(r)
doc.add_heading("8.1 Possibili metriche oggettive", level=2)
for m in ["tempo per completare il livello e numero di stelline raccolte;",
          "numero di errori (asteroidi/bombe toccati) e vite perse;",
          "ampiezza e fluidità dei movimenti rilevati dalla webcam;",
          "progressione nel tempo (confronto tra sessioni successive);",
          "tolleranza al livello di difficoltà impostato."]:
    bullet(m)
p("Questi dati, salvati a fine sessione, costituirebbero un report utile al terapista per tarare gli esercizi e "
  "seguire i progressi, e renderebbero possibile, in prospettiva, uno studio pilota di validazione.")

# ============================ 9. CONCLUSIONI ============================
doc.add_heading("9. Conclusioni e sviluppi futuri", level=1)
p("È stata realizzata un'interfaccia software completa e funzionante, scritta interamente da codice, che trasforma "
  "l'esercizio riabilitativo in un gioco e offre tre modalità di comando — incluso il movimento del braccio tramite "
  "webcam — con feedback immediato e parametri regolabili pensati per la terapia. Dal punto di vista ingegneristico il "
  "risultato è solido e portabile (Windows e Mac).")
p("Gli sviluppi naturali riguardano la misura e la validazione: salvataggio dei progressi e delle metriche, profili "
  "paziente e report per il terapista, una calibrazione guidata della webcam, nuovi livelli ed esercizi mirati a gesti "
  "specifici, e infine una valutazione con i terapisti e piccoli studi pilota per stimarne l'efficacia reale.")

# ============================ BIBLIOGRAFIA ============================
doc.add_heading("Bibliografia", level=1)
bib = LIT.get("bibliografia", [])
if bib:
    for i, b in enumerate(bib, 1):
        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(6)
        rr = par.add_run(f"[{i}] "); rr.font.bold = True; rr.font.color.rgb = TEAL_DARK
        par.add_run(b.get("citazione", ""))
else:
    p("(Da completare con i riferimenti citati.)", italic=True, color=GREY)
p("")
p("Nota metodologica: i riferimenti elencati sono stati selezionati e sottoposti a verifica di esistenza; si "
  "raccomanda comunque di ricontrollare ogni citazione e i relativi dati (autori, anno, rivista, DOI) sulla fonte "
  "originale prima della consegna.", italic=True, size=10, color=GREY)

# ============================ APPENDICE ============================
doc.add_page_break()
doc.add_heading("Appendice A — La «pagina dei parametri»", level=1)
p("Tutti i valori regolabili del gioco sono raccolti in poche pagine di parametri, separate dalla logica. Modificandoli "
  "si adattano difficoltà e accessibilità senza riscrivere il codice. Esempi:")
ptbl = doc.add_table(rows=1, cols=3); ptbl.style = "Light Grid Accent 1"
hdr = ["Parametro", "Valore tipico", "Effetto"]
for j, val in enumerate(hdr):
    c = ptbl.cell(0, j); c.text = ""; rr = c.paragraphs[0].add_run(val)
    rr.font.bold = True; rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); rr.font.size = Pt(10.5); shade(c, "006E6C")
prows = [("Vite per livello", "5", "quante volte si può sbagliare"),
         ("Tempo per livello", "60 s", "durata massima del livello"),
         ("Stelline (livello 1)", "10", "quanti bersagli raccogliere"),
         ("Morbidezza (webcam)", "regolabile", "calma il puntatore se la mano trema"),
         ("Volume errore", "più basso del positivo", "il suono di errore informa, non spaventa"),
         ("Suono / effetti visivi", "on / off", "si spengono per pazienti sensibili"),
         ("Tempo pulsanti «dwell»", "1,2 s", "quanto tenere Astro sul pulsante per premerlo")]
for r in prows:
    cells = ptbl.add_row().cells
    for j, val in enumerate(r):
        cells[j].text = ""; rr = cells[j].paragraphs[0].add_run(val); rr.font.size = Pt(10.5)
p("")
doc.add_heading("Appendice B — Galleria immagini", level=1)
p("Spazi suggeriti per le schermate del gioco da inserire nella relazione:")
image_placeholder("Figura B.1 — Schermata iniziale con la scelta del comando.")
image_placeholder("Figura B.2 — Feedback positivo: Astro si gonfia raccogliendo una stellina.")
image_placeholder("Figura B.3 — Schermata di vittoria con il pianeta amico e i coriandoli.")

doc.save(OUT)
print("Tesina salvata in", OUT, "—", len(doc.paragraphs), "paragrafi")
